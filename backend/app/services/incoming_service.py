from __future__ import annotations

import json
import sqlite3
import uuid
from datetime import UTC, date, datetime, time, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.config import DATABASE_PATH
from app.services.data_service import ensure_database_schema, normalize_description

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "request_id": ("номер запроса", "номер обращения", "id", "request_id"),
    "registration_date": ("дата регистрации", "дата", "registration_date"),
    "user_name": ("пользователь", "заявитель", "user", "user_name"),
    "service": ("услуга", "service"),
    "component": ("компонент услуги 1 уровня", "компонент", "component"),
    "request_type": ("тип запроса", "request_type"),
    "description": ("описание", "описание 2", "описание обращения", "текст обращения", "description"),
    "criticality": ("критичность", "criticality"),
    "urgency": ("срочность", "urgency"),
    "priority": ("приоритет", "priority"),
    "service_class": ("класс обслуживания", "service_class"),
    "timezone": ("часовой пояс", "часовой пояс запроса", "timezone"),
}
OPTIONAL_FIELDS = tuple(key for key in FIELD_ALIASES if key != "description")
INFERENCE_FIELDS = (
    "description", "service", "component", "request_type", "criticality", "urgency", "priority",
    "service_class", "timezone",
)
HISTORICAL_LABEL_ALIASES = (
    "вид запроса",
    "кем решен (группа)",
    "кем решён (группа)",
)
LOW_INFORMATION_DESCRIPTION_LENGTH = 10
REGISTRATION_FIELDS = (
    "request_id", "registration_date", "user_name", "service", "component", "request_type",
    "description", "criticality", "urgency", "priority", "service_class", "timezone",
)


def _connect(database: Path = DATABASE_PATH) -> sqlite3.Connection:
    ensure_database_schema(database)
    connection = sqlite3.connect(database)
    connection.row_factory = sqlite3.Row
    return connection


def _json_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    return value


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalized_header(value: Any) -> str:
    return " ".join(_text(value).lower().replace("ё", "е").split())


def _column_map(headers: list[str]) -> dict[str, int]:
    normalized = [_normalized_header(header) for header in headers]
    result: dict[str, int] = {}
    for field, aliases in FIELD_ALIASES.items():
        normalized_aliases = {_normalized_header(alias) for alias in aliases}
        for index, header in enumerate(normalized):
            if header in normalized_aliases:
                result[field] = index
                break
    return result


def _is_labeled_historical(headers: list[str]) -> bool:
    normalized_headers = {_normalized_header(header) for header in headers}
    return any(_normalized_header(alias) in normalized_headers for alias in HISTORICAL_LABEL_ALIASES)


def _has_usable_analysis_input(item: dict[str, Any]) -> bool:
    return any(_text(item.get(field)) for field in INFERENCE_FIELDS)


def _batch_payload(connection: sqlite3.Connection, batch_id: str) -> dict[str, Any] | None:
    row = connection.execute("SELECT * FROM incoming_batches WHERE id=?", (batch_id,)).fetchone()
    if row is None:
        return None
    result = dict(row)
    result["recognized_columns"] = json.loads(result.pop("recognized_columns_json"))
    result["missing_optional"] = json.loads(result.pop("missing_optional_json"))
    result["labeled_historical"] = bool(result.get("labeled_historical"))
    counts = connection.execute(
        "SELECT state, COUNT(*) count FROM incoming_tickets WHERE batch_id=? GROUP BY state", (batch_id,)
    ).fetchall()
    result["state_counts"] = {item["state"]: item["count"] for item in counts}
    result["duplicate_ids"] = [
        item[0]
        for item in connection.execute(
            "SELECT request_id FROM incoming_tickets WHERE batch_id=? AND request_id IS NOT NULL AND request_id != '' "
            "GROUP BY request_id HAVING COUNT(*) > 1 ORDER BY request_id",
            (batch_id,),
        ).fetchall()
    ]
    return result


def create_batch_from_excel(content: bytes, filename: str, database: Path = DATABASE_PATH) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xls":
        raise ValueError("Формат .xls не поддерживается текущим Excel-ридером. Сохраните файл как .xlsx.")
    if suffix != ".xlsx":
        raise ValueError("Поддерживается Excel-файл .xlsx")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("Не удалось прочитать Excel. Проверьте, что файл не повреждён.") from error
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as error:
        workbook.close()
        raise ValueError("Excel-файл пуст") from error
    last_named = max((index for index, value in enumerate(raw_headers) if _text(value)), default=-1)
    if last_named < 0:
        workbook.close()
        raise ValueError("В Excel не найдена строка заголовков")
    headers = [_text(value) for value in raw_headers[: last_named + 1]]
    mapping = _column_map(headers)

    batch_id = uuid.uuid4().hex
    created_at = datetime.now(UTC).isoformat()
    recognized_columns = [headers[index] for index in sorted(set(mapping.values()))]
    missing_optional = [field for field in OPTIONAL_FIELDS if field not in mapping]
    description_column = headers[mapping["description"]] if "description" in mapping else None
    labeled_historical = _is_labeled_historical(headers)
    parsed: list[dict[str, Any]] = []
    total_rows = valid_rows = invalid_rows = 0
    for row_number, row in enumerate(rows, start=2):
        values = list(row[: len(headers)])
        if not any(value is not None and _text(value) for value in values):
            continue
        total_rows += 1
        raw = {header: _json_value(values[index] if index < len(values) else None) for index, header in enumerate(headers)}
        item = {field: _text(values[index] if index < len(values) else None) for field, index in mapping.items()}
        description = item.get("description", "")
        low_information = len(description) < LOW_INFORMATION_DESCRIPTION_LENGTH
        error_text = "" if _has_usable_analysis_input(item) else "Недостаточно данных для анализа: нет текста обращения и регистрационных признаков"
        state = "uploaded" if not error_text else "failed"
        valid_rows += int(not error_text)
        invalid_rows += int(bool(error_text))
        parsed.append({
            **{field: item.get(field, "") for field in FIELD_ALIASES},
            "row_number": row_number,
            "raw_json": json.dumps(raw, ensure_ascii=False, sort_keys=True),
            "state": state,
            "error_text": error_text or None,
            "low_information": int(low_information),
        })
    workbook.close()
    if total_rows == 0:
        raise ValueError("В Excel нет строк с обращениями")

    connection = _connect(database)
    connection.execute(
        """
        INSERT INTO incoming_batches (
            id, filename, status, total_rows, valid_rows, invalid_rows,
            recognized_columns_json, missing_optional_json, description_column,
            labeled_historical, created_at, analyzed_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL)
        """,
        (
            batch_id, Path(filename).name, "uploaded", total_rows, valid_rows, invalid_rows,
            json.dumps(recognized_columns, ensure_ascii=False),
            json.dumps(missing_optional, ensure_ascii=False), description_column,
            int(labeled_historical), created_at,
        ),
    )
    for item in parsed:
        connection.execute(
            """
            INSERT INTO incoming_tickets (
                batch_id, row_number, request_id, registration_date, user_name, service, component,
                request_type, description, criticality, urgency, priority, service_class, timezone,
                raw_json, state, error_text, low_information
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                batch_id, item["row_number"], item["request_id"] or None, item["registration_date"] or None,
                item["user_name"] or None, item["service"] or None, item["component"] or None,
                item["request_type"] or None, item["description"], item["criticality"] or None,
                item["urgency"] or None, item["priority"] or None, item["service_class"] or None,
                item["timezone"] or None, item["raw_json"], item["state"], item["error_text"],
                item["low_information"],
            ),
        )
    connection.commit()
    result = _batch_payload(connection, batch_id)
    connection.close()
    assert result is not None
    return result


def create_manual_ticket(payload: dict[str, Any], database: Path = DATABASE_PATH) -> dict[str, Any]:
    canonical = {
        "description": _text(payload.get("description")),
        "service": _text(payload.get("service")),
        "component": _text(payload.get("component")),
        "request_type": _text(payload.get("request_type")),
        "criticality": _text(payload.get("criticality")),
        "urgency": _text(payload.get("urgency")),
        "priority": _text(payload.get("priority")),
        "service_class": _text(payload.get("service_class")),
        "timezone": _text(payload.get("timezone")),
    }
    if not _has_usable_analysis_input(canonical):
        raise ValueError("Недостаточно данных для анализа: укажите текст обращения или регистрационные признаки")
    low_information = int(len(canonical["description"]) < LOW_INFORMATION_DESCRIPTION_LENGTH)
    batch_id = uuid.uuid4().hex
    created_at = datetime.now(UTC).isoformat()
    connection = _connect(database)
    connection.execute(
        """
        INSERT INTO incoming_batches (
            id, filename, status, total_rows, valid_rows, invalid_rows,
            recognized_columns_json, missing_optional_json, description_column,
            labeled_historical, created_at, analyzed_at
        ) VALUES (?, ?, ?, 1, 1, 0, ?, '[]', ?, 0, ?, NULL)
        """,
        (
            batch_id, "Одно обращение", "uploaded",
            json.dumps([key for key, value in canonical.items() if _text(value)], ensure_ascii=False),
            "Описание" if canonical["description"] else None,
            created_at,
        ),
    )
    cursor = connection.execute(
        """
        INSERT INTO incoming_tickets (
            batch_id, row_number, request_id, registration_date, user_name, service, component,
            request_type, description, criticality, urgency, priority, service_class, timezone,
            raw_json, state, low_information
        ) VALUES (?, 1, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'uploaded', ?)
        """,
        (
            batch_id, payload.get("request_id") or None, payload.get("registration_date") or None,
            payload.get("user") or None, payload.get("service") or None, payload.get("component") or None,
            payload.get("request_type") or None, canonical["description"], payload.get("criticality") or None,
            payload.get("urgency") or None, payload.get("priority") or None, payload.get("service_class") or None,
            payload.get("timezone") or None, json.dumps(payload, ensure_ascii=False, sort_keys=True), low_information,
        ),
    )
    connection.commit()
    result = _batch_payload(connection, batch_id) or {}
    result["ticket_id"] = cursor.lastrowid
    connection.close()
    return result


def list_batches(database: Path = DATABASE_PATH) -> list[dict[str, Any]]:
    connection = _connect(database)
    ids = [row[0] for row in connection.execute("SELECT id FROM incoming_batches ORDER BY created_at DESC LIMIT 25")]
    result = [_batch_payload(connection, batch_id) for batch_id in ids]
    connection.close()
    return [item for item in result if item is not None]


def get_batch(batch_id: str, database: Path = DATABASE_PATH) -> dict[str, Any] | None:
    connection = _connect(database)
    result = _batch_payload(connection, batch_id)
    connection.close()
    return result


def _ticket_payload(row: sqlite3.Row) -> dict[str, Any]:
    item = dict(row)
    item["low_information"] = bool(item.get("low_information"))
    item["needs_training_review"] = bool(item.get("needs_training_review"))
    item["raw"] = json.loads(item.pop("raw_json") or "{}")
    analysis_json = item.pop("analysis_json")
    item["analysis"] = json.loads(analysis_json) if analysis_json else None
    item["attention"] = bool(
        item["state"] == "failed"
        or item["low_information"]
        or (
            item["analysis"]
            and (
                item["analysis"]["category"].get("review_required")
                or item["analysis"]["routing"].get("review_required")
            )
        )
    )
    return item


def list_batch_tickets(batch_id: str, database: Path = DATABASE_PATH) -> list[dict[str, Any]]:
    connection = _connect(database)
    rows = connection.execute("SELECT * FROM incoming_tickets WHERE batch_id=? ORDER BY id", (batch_id,)).fetchall()
    result = [_ticket_payload(row) for row in rows]
    connection.close()
    return result


def get_incoming_ticket(ticket_id: int, database: Path = DATABASE_PATH) -> dict[str, Any] | None:
    connection = _connect(database)
    row = connection.execute("SELECT * FROM incoming_tickets WHERE id=?", (ticket_id,)).fetchone()
    result = _ticket_payload(row) if row else None
    connection.close()
    return result


def _audit_snapshot(
    connection: sqlite3.Connection,
    row: sqlite3.Row,
    event_type: str,
    extra: dict[str, Any] | None = None,
) -> None:
    analysis = json.loads(row["analysis_json"]) if row["analysis_json"] else None
    snapshot: dict[str, Any] = {
        "registration": {field: row[field] for field in REGISTRATION_FIELDS},
        "analysis": analysis,
        "model_category": row["model_category"],
        "model_category_confidence": row["model_category_confidence"],
        "model_route": row["model_route"],
        "model_route_confidence": row["model_route_confidence"],
        "operator_category": row["operator_category"],
        "operator_route": row["operator_route"],
        "category_confirmed": bool(row["category_confirmed"]),
        "route_confirmed": bool(row["route_confirmed"]),
    }
    if extra:
        snapshot.update(extra)
    connection.execute(
        "INSERT INTO incoming_ticket_audit (incoming_ticket_id, event_type, snapshot_json, created_at) "
        "VALUES (?, ?, ?, ?)",
        (row["id"], event_type, json.dumps(snapshot, ensure_ascii=False, sort_keys=True), datetime.now(UTC).isoformat()),
    )


def update_incoming_ticket(
    ticket_id: int,
    payload: dict[str, Any],
    database: Path = DATABASE_PATH,
) -> dict[str, Any] | None:
    connection = _connect(database)
    row = connection.execute("SELECT * FROM incoming_tickets WHERE id=?", (ticket_id,)).fetchone()
    if row is None:
        connection.close()
        return None
    if row["state"] == "saved":
        connection.close()
        raise ValueError("Сохранённое обращение нельзя изменять через очередь обработки")

    values = {field: row[field] for field in REGISTRATION_FIELDS}
    key_map = {"user": "user_name"}
    for key, value in payload.items():
        field = key_map.get(key, key)
        if field not in values:
            continue
        normalized = _text(value)
        values[field] = (normalized or None) if field in {"request_id", "registration_date"} else normalized
    values["description"] = _text(values.get("description"))
    if not _has_usable_analysis_input(values):
        connection.close()
        raise ValueError("Недостаточно данных для анализа: укажите текст обращения или регистрационные признаки")

    request_id = values["request_id"]
    if request_id and request_id != row["request_id"]:
        duplicate_staging = connection.execute(
            "SELECT 1 FROM incoming_tickets WHERE request_id=? AND id<>? LIMIT 1",
            (request_id, ticket_id),
        ).fetchone()
        duplicate_saved = connection.execute(
            "SELECT 1 FROM tickets WHERE request_id=? LIMIT 1", (request_id,)
        ).fetchone()
        if duplicate_staging or duplicate_saved:
            connection.close()
            raise ValueError(f"Номер обращения {request_id} уже используется")

    changed = any((row[field] or "") != (values[field] or "") for field in REGISTRATION_FIELDS)
    if not changed:
        result = _ticket_payload(row)
        connection.close()
        return result

    connection.execute("BEGIN")
    try:
        _audit_snapshot(connection, row, "source_edit")
        connection.execute(
            """
            UPDATE incoming_tickets
            SET request_id=?, registration_date=?, user_name=?, service=?, component=?, request_type=?,
                description=?, criticality=?, urgency=?, priority=?, service_class=?, timezone=?,
                state='uploaded', analysis_json=NULL, model_category=NULL, model_category_confidence=NULL,
                model_route=NULL, model_route_confidence=NULL, operator_category=NULL, operator_route=NULL,
                category_confirmed=0, route_confirmed=0, reviewed_at=NULL, saved_request_id=NULL,
                error_text=NULL, low_information=?, prediction_timestamp=NULL,
                learning_priority=0, needs_training_review=0
            WHERE id=?
            """,
            (
                values["request_id"], values["registration_date"], values["user_name"] or None,
                values["service"] or None, values["component"] or None, values["request_type"] or None,
                values["description"], values["criticality"] or None, values["urgency"] or None,
                values["priority"] or None, values["service_class"] or None, values["timezone"] or None,
                int(len(values["description"]) < LOW_INFORMATION_DESCRIPTION_LENGTH), ticket_id,
            ),
        )
        connection.execute(
            "UPDATE incoming_batches SET status='waiting_review' WHERE id=?",
            (row["batch_id"],),
        )
        connection.commit()
    except Exception:
        connection.rollback()
        connection.close()
        raise
    refreshed = connection.execute("SELECT * FROM incoming_tickets WHERE id=?", (ticket_id,)).fetchone()
    result = _ticket_payload(refreshed)
    connection.close()
    return result


def ticket_analysis_payload(ticket_id: int, database: Path = DATABASE_PATH) -> dict[str, str] | None:
    ticket = get_incoming_ticket(ticket_id, database)
    if ticket is None:
        return None
    return {
        "description": ticket["description"],
        "service": ticket.get("service") or "",
        "component": ticket.get("component") or "",
        "request_type": ticket.get("request_type") or "",
        "criticality": ticket.get("criticality") or "",
        "urgency": ticket.get("urgency") or "",
        "priority": ticket.get("priority") or "",
        "service_class": ticket.get("service_class") or "",
        "timezone": ticket.get("timezone") or "",
        "registration_date": ticket.get("registration_date") or "",
        "user": ticket.get("user_name") or "",
    }


def set_batch_status(batch_id: str, status: str, database: Path = DATABASE_PATH) -> None:
    connection = _connect(database)
    connection.execute("UPDATE incoming_batches SET status=? WHERE id=?", (status, batch_id))
    connection.commit()
    connection.close()

# staging workflow


def _learning_signal(result: dict[str, Any], *, low_information: bool = False) -> tuple[float, int]:
    category = result.get("category") or {}
    routing = result.get("routing") or {}
    alternatives = category.get("alternatives") or []
    priority = 0.0
    if category.get("review_required"):
        priority += 0.4
    if routing.get("review_required"):
        priority += 0.25
    if len(alternatives) > 1:
        margin = float(alternatives[0].get("confidence") or 0.0) - float(alternatives[1].get("confidence") or 0.0)
        if margin < 0.1:
            priority += 0.2
    if low_information:
        priority += 0.15
    priority = round(min(priority, 1.0), 3)
    return priority, int(priority >= 0.25)


def save_analysis(ticket_id: int, result: dict[str, Any], database: Path = DATABASE_PATH) -> None:
    connection = _connect(database)
    before = connection.execute("SELECT * FROM incoming_tickets WHERE id=?", (ticket_id,)).fetchone()
    if before is None:
        connection.close()
        return
    if before["state"] == "saved":
        connection.close()
        raise ValueError("Сохранённое обращение нельзя анализировать повторно")
    learning_priority, needs_training_review = _learning_signal(
        result, low_information=bool(before["low_information"])
    )
    prediction_timestamp = datetime.now(UTC).isoformat()
    sql = (
        "UPDATE incoming_tickets SET state=?, analysis_json=?, model_category=?, "
        "model_category_confidence=?, model_route=?, model_route_confidence=?, error_text=NULL, "
        "prediction_timestamp=?, learning_priority=?, needs_training_review=? WHERE id=?"
    )
    values = (
        "waiting_review",
        json.dumps(result, ensure_ascii=False),
        result["category"]["label"],
        result["category"]["confidence"],
        result["routing"]["label"],
        result["routing"]["confidence"],
        prediction_timestamp,
        learning_priority,
        needs_training_review,
        ticket_id,
    )
    connection.execute(sql, values)
    refreshed = connection.execute("SELECT * FROM incoming_tickets WHERE id=?", (ticket_id,)).fetchone()
    _audit_snapshot(connection, refreshed, "analysis")
    connection.commit()
    connection.close()


def save_analysis_error(ticket_id: int, message: str, database: Path = DATABASE_PATH) -> None:
    connection = _connect(database)
    before = connection.execute("SELECT * FROM incoming_tickets WHERE id=?", (ticket_id,)).fetchone()
    if before is None:
        connection.close()
        return
    connection.execute(
        "UPDATE incoming_tickets SET state=?, error_text=? WHERE id=?",
        ("failed", message, ticket_id),
    )
    refreshed = connection.execute("SELECT * FROM incoming_tickets WHERE id=?", (ticket_id,)).fetchone()
    _audit_snapshot(connection, refreshed, "analysis_error", {"error": message})
    connection.commit()
    connection.close()


def finish_batch_analysis(batch_id: str, database: Path = DATABASE_PATH) -> dict[str, Any] | None:
    connection = _connect(database)
    connection.execute(
        "UPDATE incoming_batches SET status=?, analyzed_at=? WHERE id=?",
        ("waiting_review", datetime.now(UTC).isoformat(), batch_id),
    )
    connection.commit()
    result = _batch_payload(connection, batch_id)
    connection.close()
    return result

# operator decisions


def confirm_category(ticket_id: int, category: str, database: Path = DATABASE_PATH) -> dict[str, Any] | None:
    return _confirm_choice(ticket_id, "category", category, database)


def confirm_route(ticket_id: int, route: str, database: Path = DATABASE_PATH) -> dict[str, Any] | None:
    return _confirm_choice(ticket_id, "route", route, database)


def _confirm_choice(ticket_id: int, kind: str, value: str, database: Path = DATABASE_PATH) -> dict[str, Any] | None:
    if kind not in {"category", "route"}:
        raise ValueError("Unsupported decision")
    connection = _connect(database)
    row = connection.execute("SELECT * FROM incoming_tickets WHERE id=?", (ticket_id,)).fetchone()
    if row is None:
        connection.close()
        return None
    if row["state"] == "saved":
        connection.close()
        raise ValueError("Обращение уже сохранено")
    if not row["analysis_json"]:
        connection.close()
        raise ValueError("Сначала выполните анализ обращения")
    value = value.strip()
    if not value:
        connection.close()
        raise ValueError("Значение решения не может быть пустым")
    if kind == "category":
        connection.execute(
            "UPDATE incoming_tickets SET operator_category=?, category_confirmed=1, "
            "operator_route=NULL, route_confirmed=0 WHERE id=?",
            (value, ticket_id),
        )
    else:
        connection.execute(
            "UPDATE incoming_tickets SET operator_route=?, route_confirmed=1 WHERE id=?",
            (value, ticket_id),
        )
    refreshed = connection.execute("SELECT * FROM incoming_tickets WHERE id=?", (ticket_id,)).fetchone()
    _audit_snapshot(
        connection,
        refreshed,
        "category_confirmed" if kind == "category" else "route_confirmed",
        {"selected_value": value},
    )
    connection.commit()
    result = _ticket_payload(refreshed)
    connection.close()
    return result


def save_routing_recommendation(
    ticket_id: int,
    routing: dict[str, Any],
    database: Path = DATABASE_PATH,
) -> dict[str, Any] | None:
    connection = _connect(database)
    row = connection.execute("SELECT * FROM incoming_tickets WHERE id=?", (ticket_id,)).fetchone()
    if row is None:
        connection.close()
        return None
    if not row["analysis_json"]:
        connection.close()
        raise ValueError("Сначала выполните анализ обращения")
    analysis = json.loads(row["analysis_json"])
    analysis["routing"] = routing
    connection.execute(
        "UPDATE incoming_tickets SET analysis_json=?, model_route=?, model_route_confidence=?, "
        "operator_route=NULL, route_confirmed=0 WHERE id=?",
        (
            json.dumps(analysis, ensure_ascii=False),
            routing.get("label"),
            routing.get("confidence"),
            ticket_id,
        ),
    )
    refreshed = connection.execute("SELECT * FROM incoming_tickets WHERE id=?", (ticket_id,)).fetchone()
    _audit_snapshot(connection, refreshed, "routing_recalculated", {"routing_mode": routing.get("mode")})
    connection.commit()
    result = _ticket_payload(refreshed)
    connection.close()
    return result


def _unique_saved_request_id(connection: sqlite3.Connection, ticket: sqlite3.Row) -> str:
    source_id = _text(ticket["request_id"])
    if source_id:
        exists = connection.execute("SELECT 1 FROM tickets WHERE request_id=?", (source_id,)).fetchone()
        if exists is None:
            return source_id
    base = f"NEW-{ticket['batch_id'][:8]}-{ticket['id']}"
    candidate = base
    index = 2
    while connection.execute("SELECT 1 FROM tickets WHERE request_id=?", (candidate,)).fetchone():
        candidate = f"{base}-{index}"
        index += 1
    return candidate


def complete_ticket(ticket_id: int, database: Path = DATABASE_PATH) -> dict[str, Any] | None:
    connection = _connect(database)
    row = connection.execute("SELECT * FROM incoming_tickets WHERE id=?", (ticket_id,)).fetchone()
    if row is None:
        connection.close()
        return None
    if row["state"] == "saved":
        saved_id = row["saved_request_id"]
        connection.close()
        return {"status": "saved", "request_id": saved_id, "incoming_ticket_id": ticket_id}
    if not row["analysis_json"]:
        connection.close()
        raise ValueError("Сначала выполните анализ обращения")
    if not row["category_confirmed"] or not row["route_confirmed"]:
        connection.close()
        raise ValueError("Сначала подтвердите категорию и линию поддержки")

    analysis = json.loads(row["analysis_json"])
    final_category = row["operator_category"] or row["model_category"]
    final_route = row["operator_route"] or row["model_route"]
    if not final_category or not final_route:
        connection.close()
        raise ValueError("Не удалось определить итоговую категорию или линию")

    saved_request_id = _unique_saved_request_id(connection, row)
    now = datetime.now(UTC).isoformat()
    registration_date = row["registration_date"] or now
    raw = json.loads(row["raw_json"] or "{}")
    raw["_operator_review"] = {
        "model_category": row["model_category"],
        "model_category_confidence": row["model_category_confidence"],
        "model_top3": analysis.get("category", {}).get("alternatives", [])[:3],
        "model_route": row["model_route"],
        "model_route_confidence": row["model_route_confidence"],
        "operator_final_category": final_category,
        "operator_final_route": final_route,
        "reviewed_at": now,
    }
    provenance = analysis.get("model_provenance") or {}
    model_version = (
        provenance.get("version")
        or provenance.get("candidate_id")
        or provenance.get("family")
        or provenance.get("policy")
    )
    category_changed = int(final_category != row["model_category"])
    route_changed = int(final_route != row["model_route"])
    accepted_without_change = int(not category_changed and not route_changed)
    registration = {field: row[field] for field in REGISTRATION_FIELDS}
    available_fields = [field for field, value in registration.items() if _text(value)]
    missing_fields = [field for field, value in registration.items() if not _text(value)]
    learning_priority = float(row["learning_priority"] or 0.0)
    if category_changed or route_changed:
        learning_priority = max(learning_priority, 1.0)
    needs_training_review = int(bool(row["needs_training_review"]) or category_changed or route_changed)

    connection.execute("BEGIN")
    try:
        connection.execute(
            """
            INSERT INTO tickets (
                request_id, registration_date, user_name, service, component, category, request_type, description,
                normalized_description, criticality, urgency, priority, service_class, timezone, status,
                overdue, final_line, actual_duration_seconds, clarifications_count, clarification_seconds,
                line1_participated, line2_participated, line3_participated, line4_participated, participants_count,
                result, raw_json, source_hash, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, ?, NULL, 0, NULL, 0, 0, 0, 0, 0, ?, ?, ?, ?)
            """,
            (
                saved_request_id,
                registration_date,
                row["user_name"] or None,
                row["service"] or "",
                row["component"] or None,
                final_category,
                row["request_type"] or "",
                row["description"],
                normalize_description(row["description"]),
                row["criticality"] or "",
                row["urgency"] or "",
                row["priority"] or "",
                row["service_class"] or "",
                row["timezone"] or "",
                "Проверено оператором",
                final_route,
                "Проверено оператором",
                json.dumps(raw, ensure_ascii=False, sort_keys=True),
                "operator-review",
                now,
            ),
        )
        connection.execute(
            """
            INSERT INTO ticket_reviews (
                incoming_ticket_id, batch_id, source_request_id, saved_request_id,
                model_category, model_category_confidence, model_top3_json,
                model_route, model_route_confidence, operator_final_category, operator_final_route,
                category_changed, route_changed, model_version, reviewed_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                ticket_id, row["batch_id"], row["request_id"], saved_request_id,
                row["model_category"], row["model_category_confidence"],
                json.dumps(analysis.get("category", {}).get("alternatives", [])[:3], ensure_ascii=False),
                row["model_route"], row["model_route_confidence"], final_category, final_route,
                category_changed, route_changed, model_version, now,
            ),
        )
        connection.execute(
            """
            INSERT INTO production_feedback (
                incoming_ticket_id, request_id, registration_json,
                model_category, model_category_confidence, model_top3_json,
                model_route, model_route_confidence,
                operator_final_category, operator_final_route,
                category_corrected, route_corrected, accepted_without_change,
                model_version, prediction_timestamp, confirmation_timestamp,
                available_fields_json, missing_fields_json,
                learning_priority, needs_training_review
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                ticket_id,
                saved_request_id,
                json.dumps(registration, ensure_ascii=False, sort_keys=True),
                row["model_category"],
                row["model_category_confidence"],
                json.dumps(analysis.get("category", {}).get("alternatives", [])[:3], ensure_ascii=False),
                row["model_route"],
                row["model_route_confidence"],
                final_category,
                final_route,
                category_changed,
                route_changed,
                accepted_without_change,
                model_version,
                row["prediction_timestamp"],
                now,
                json.dumps(available_fields, ensure_ascii=False),
                json.dumps(missing_fields, ensure_ascii=False),
                learning_priority,
                needs_training_review,
            ),
        )
        connection.execute(
            "UPDATE incoming_tickets SET state='saved', reviewed_at=?, saved_request_id=? WHERE id=?",
            (now, saved_request_id, ticket_id),
        )
        refreshed = connection.execute("SELECT * FROM incoming_tickets WHERE id=?", (ticket_id,)).fetchone()
        _audit_snapshot(
            connection,
            refreshed,
            "final_approved",
            {"final_category": final_category, "final_route": final_route, "saved_request_id": saved_request_id},
        )
        remaining = connection.execute(
            "SELECT COUNT(*) FROM incoming_tickets WHERE batch_id=? AND state NOT IN ('saved', 'failed')",
            (row["batch_id"],),
        ).fetchone()[0]
        connection.execute(
            "UPDATE incoming_batches SET status=? WHERE id=?",
            ("saved" if remaining == 0 else "waiting_review", row["batch_id"]),
        )
        connection.commit()
    except Exception:
        connection.rollback()
        connection.close()
        raise
    connection.close()
    return {"status": "saved", "request_id": saved_request_id, "incoming_ticket_id": ticket_id}
