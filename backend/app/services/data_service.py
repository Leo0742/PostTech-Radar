from __future__ import annotations

import hashlib
import json
import math
import re
import sqlite3
from collections import Counter
from datetime import UTC, date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook

from app.core.config import DATABASE_PATH, RAW_DATA_DIR, ensure_directories
from app.ml.v3.dataset import (
    DatasetConfig,
    challenge_dataset_config,
    source_value,
    training_corpus_summary,
)

_DURATION_RE = re.compile(r"^(\d+):([0-5]\d):([0-5]\d)$")
SCHEMA_VERSION = "7"


def normalize_description(value: Any) -> str:
    return " ".join(str(value or "").lower().split())


def parse_duration(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, timedelta):
        return int(value.total_seconds())
    if isinstance(value, time):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, (int, float)):
        return int(value * 86_400) if 0 <= value < 10_000 else int(value)
    match = _DURATION_RE.fullmatch(str(value).strip())
    if not match:
        return None
    hours, minutes, seconds = (int(part) for part in match.groups())
    return hours * 3600 + minutes * 60 + seconds


def human_duration(seconds: int | None) -> str | None:
    if seconds is None:
        return None
    hours, remainder = divmod(max(0, int(seconds)), 3600)
    minutes = remainder // 60
    return f"{hours} ч {minutes} мин"


def participation_flags(row: dict[str, Any]) -> tuple[int, int, int, int]:
    flags: list[int] = []
    for line in range(1, 5):
        reaction = parse_duration(row.get(f"Суммарное время реакции {line} линии")) or 0
        work = parse_duration(row.get(f"Суммарное время работы {line} линии")) or 0
        flags.append(int(reaction > 0 or work > 0))
    return tuple(flags)  # type: ignore[return-value]


def _json_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        total = int(value.total_seconds())
        return f"{total // 3600:02d}:{(total % 3600) // 60:02d}:{total % 60:02d}"
    return value


def _request_id(value: Any) -> str:
    if value is None or str(value).strip() == "":
        raise ValueError("Номер запроса не может быть пустым")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def find_dataset(path: Path | None = None) -> Path:
    if path:
        resolved = path.expanduser().resolve()
        if not resolved.is_file() or resolved.suffix.lower() != ".xlsx":
            raise FileNotFoundError(f"Не найден Excel-файл: {resolved}")
        return resolved
    candidates = sorted(RAW_DATA_DIR.glob("*.xlsx"))
    if not candidates:
        raise FileNotFoundError(f"Поместите набор данных .xlsx в {RAW_DATA_DIR}")
    return candidates[0]


def read_workbook(path: Path, dataset_config: DatasetConfig | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    dataset_config = dataset_config or challenge_dataset_config()
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    raw_headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    last_named = max(index for index, value in enumerate(raw_headers) if value)
    headers = raw_headers[: last_named + 1]
    required = dataset_config.required_columns or dataset_config.essential_columns
    missing = [column for column in required if column not in headers]
    if missing:
        raise ValueError(f"В Excel отсутствуют обязательные колонки: {', '.join(missing)}")
    records = [
        dict(zip(headers, row[: len(headers)], strict=True))
        for row in rows
        if any(value is not None for value in row[: len(headers)])
    ]
    workbook.close()
    return headers, records


def inspect_workbook(path: Path, dataset_config: DatasetConfig | None = None) -> dict[str, Any]:
    dataset_config = dataset_config or challenge_dataset_config()
    headers, records = read_workbook(path, dataset_config)
    category_counts = Counter(row[dataset_config.label_column] for row in records)
    top15 = category_counts.most_common(dataset_config.top_k)
    return {
        "source": str(path),
        "records": len(records),
        "columns": len(headers),
        "distinct_categories": len(category_counts),
        "top15_records": sum(count for _, count in top15),
        "top15": [{"name": str(name), "count": count} for name, count in top15],
        "support_lines": dict(Counter(row[dataset_config.route_column] for row in records)),
        "sla_labels": dict(Counter(row.get(dataset_config.operational_columns.get("overdue", "")) for row in records)),
        "statuses": dict(Counter(row.get(dataset_config.operational_columns.get("status", "")) for row in records)),
        "missing": {header: sum(row[header] is None for row in records) for header in headers},
    }


def _connect(path: Path = DATABASE_PATH) -> sqlite3.Connection:
    connection = sqlite3.connect(path)
    connection.row_factory = sqlite3.Row
    return connection


_TICKET_COLUMNS = (
    "request_id", "registration_date", "user_name", "service", "component", "category", "request_type", "description",
    "normalized_description", "criticality", "urgency", "priority", "service_class", "timezone", "status",
    "overdue", "final_line", "actual_duration_seconds", "clarifications_count", "clarification_seconds",
    "line1_participated", "line2_participated", "line3_participated", "line4_participated", "participants_count",
    "result", "raw_json", "source_hash", "updated_at",
)


def _create_tickets_table(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS tickets (
            request_id TEXT PRIMARY KEY,
            registration_date TEXT NOT NULL,
            user_name TEXT,
            service TEXT NOT NULL,
            component TEXT,
            category TEXT NOT NULL,
            request_type TEXT NOT NULL,
            description TEXT NOT NULL,
            normalized_description TEXT NOT NULL,
            criticality TEXT NOT NULL,
            urgency TEXT NOT NULL,
            priority TEXT NOT NULL,
            service_class TEXT NOT NULL,
            timezone TEXT NOT NULL,
            status TEXT NOT NULL,
            overdue INTEGER NOT NULL,
            final_line TEXT NOT NULL,
            actual_duration_seconds INTEGER,
            clarifications_count INTEGER NOT NULL,
            clarification_seconds INTEGER,
            line1_participated INTEGER NOT NULL DEFAULT 0,
            line2_participated INTEGER NOT NULL DEFAULT 0,
            line3_participated INTEGER NOT NULL DEFAULT 0,
            line4_participated INTEGER NOT NULL DEFAULT 0,
            participants_count INTEGER NOT NULL,
            result TEXT NOT NULL,
            raw_json TEXT NOT NULL,
            source_hash TEXT NOT NULL DEFAULT '',
            updated_at TEXT NOT NULL DEFAULT ''
        );
        CREATE INDEX IF NOT EXISTS idx_tickets_filters
            ON tickets(registration_date, service, category, priority, final_line);
        CREATE INDEX IF NOT EXISTS idx_tickets_description ON tickets(normalized_description);
        """
    )


def _legacy_values(row: sqlite3.Row) -> tuple[Any, ...]:
    item = dict(row)
    raw = json.loads(item.get("raw_json") or "{}")
    flags = participation_flags(raw)
    return (
        _request_id(item["request_id"]), item["registration_date"], item.get("user_name"), item["service"], item.get("component"),
        item["category"], item["request_type"], item["description"], item["normalized_description"],
        item["criticality"], item["urgency"], item["priority"], item["service_class"], item["timezone"],
        item["status"], item["overdue"], item["final_line"], item.get("actual_duration_seconds"),
        item["clarifications_count"], item.get("clarification_seconds"), *flags, sum(flags), item["result"],
        item["raw_json"], "legacy-migration", datetime.now(UTC).isoformat(),
    )


def ensure_database_schema(database: Path = DATABASE_PATH) -> None:
    database.parent.mkdir(parents=True, exist_ok=True)
    connection = _connect(database)
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS dataset_imports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_hash TEXT NOT NULL,
            source_path TEXT NOT NULL,
            mode TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            inserted INTEGER NOT NULL,
            updated INTEGER NOT NULL,
            unchanged INTEGER NOT NULL,
            total INTEGER NOT NULL
        );
        CREATE TABLE IF NOT EXISTS status_events (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            request_id TEXT NOT NULL,
            status TEXT NOT NULL,
            entered_at TEXT NOT NULL,
            left_at TEXT,
            source TEXT,
            UNIQUE(request_id, status, entered_at)
        );
        CREATE TABLE IF NOT EXISTS incoming_batches (
            id TEXT PRIMARY KEY,
            filename TEXT NOT NULL,
            status TEXT NOT NULL,
            total_rows INTEGER NOT NULL,
            valid_rows INTEGER NOT NULL,
            invalid_rows INTEGER NOT NULL,
            recognized_columns_json TEXT NOT NULL,
            missing_optional_json TEXT NOT NULL,
            description_column TEXT,
            labeled_historical INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            analyzed_at TEXT
        );
        CREATE TABLE IF NOT EXISTS incoming_tickets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            batch_id TEXT NOT NULL,
            row_number INTEGER NOT NULL,
            request_id TEXT,
            registration_date TEXT,
            user_name TEXT,
            service TEXT,
            component TEXT,
            request_type TEXT,
            description TEXT NOT NULL,
            criticality TEXT,
            urgency TEXT,
            priority TEXT,
            service_class TEXT,
            timezone TEXT,
            raw_json TEXT NOT NULL,
            state TEXT NOT NULL,
            analysis_json TEXT,
            model_category TEXT,
            model_category_confidence REAL,
            model_route TEXT,
            model_route_confidence REAL,
            operator_category TEXT,
            operator_route TEXT,
            category_confirmed INTEGER NOT NULL DEFAULT 0,
            route_confirmed INTEGER NOT NULL DEFAULT 0,
            reviewed_at TEXT,
            saved_request_id TEXT,
            error_text TEXT,
            low_information INTEGER NOT NULL DEFAULT 0,
            prediction_timestamp TEXT,
            learning_priority REAL NOT NULL DEFAULT 0,
            needs_training_review INTEGER NOT NULL DEFAULT 0,
            UNIQUE(batch_id, row_number),
            FOREIGN KEY(batch_id) REFERENCES incoming_batches(id)
        );
        CREATE INDEX IF NOT EXISTS idx_incoming_tickets_batch_state
            ON incoming_tickets(batch_id, state, id);
        CREATE TABLE IF NOT EXISTS ticket_reviews (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            incoming_ticket_id INTEGER NOT NULL,
            batch_id TEXT NOT NULL,
            source_request_id TEXT,
            saved_request_id TEXT NOT NULL,
            model_category TEXT,
            model_category_confidence REAL,
            model_top3_json TEXT NOT NULL,
            model_route TEXT,
            model_route_confidence REAL,
            operator_final_category TEXT NOT NULL,
            operator_final_route TEXT NOT NULL,
            category_changed INTEGER NOT NULL,
            route_changed INTEGER NOT NULL,
            model_version TEXT,
            reviewed_at TEXT NOT NULL,
            UNIQUE(incoming_ticket_id),
            FOREIGN KEY(incoming_ticket_id) REFERENCES incoming_tickets(id)
        );
        CREATE TABLE IF NOT EXISTS incoming_ticket_audit (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            incoming_ticket_id INTEGER NOT NULL,
            event_type TEXT NOT NULL,
            snapshot_json TEXT NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY(incoming_ticket_id) REFERENCES incoming_tickets(id)
        );
        CREATE INDEX IF NOT EXISTS idx_incoming_ticket_audit_ticket
            ON incoming_ticket_audit(incoming_ticket_id, id);
        CREATE TABLE IF NOT EXISTS ticket_revisions (
            revision_id INTEGER PRIMARY KEY AUTOINCREMENT,
            request_id TEXT NOT NULL,
            changed_at TEXT NOT NULL,
            changed_by TEXT NOT NULL,
            changed_fields_json TEXT NOT NULL,
            before_json TEXT NOT NULL,
            after_json TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_ticket_revisions_request
            ON ticket_revisions(request_id, revision_id DESC);
        CREATE TABLE IF NOT EXISTS production_feedback (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            incoming_ticket_id INTEGER NOT NULL UNIQUE,
            request_id TEXT NOT NULL,
            registration_json TEXT NOT NULL,
            model_category TEXT,
            model_category_confidence REAL,
            model_top3_json TEXT NOT NULL,
            model_route TEXT,
            model_route_confidence REAL,
            operator_final_category TEXT NOT NULL,
            operator_final_route TEXT NOT NULL,
            category_corrected INTEGER NOT NULL,
            route_corrected INTEGER NOT NULL,
            accepted_without_change INTEGER NOT NULL,
            model_version TEXT,
            prediction_timestamp TEXT,
            confirmation_timestamp TEXT NOT NULL,
            available_fields_json TEXT NOT NULL,
            missing_fields_json TEXT NOT NULL,
            learning_priority REAL NOT NULL DEFAULT 0,
            needs_training_review INTEGER NOT NULL DEFAULT 0,
            operator_comment TEXT,
            FOREIGN KEY(incoming_ticket_id) REFERENCES incoming_tickets(id)
        );
        CREATE INDEX IF NOT EXISTS idx_production_feedback_training_review
            ON production_feedback(needs_training_review, learning_priority DESC, id);
        """
    )
    incoming_batch_columns = {row["name"] for row in connection.execute("PRAGMA table_info(incoming_batches)")}
    if "description_column" not in incoming_batch_columns:
        connection.execute("ALTER TABLE incoming_batches ADD COLUMN description_column TEXT")
    if "labeled_historical" not in incoming_batch_columns:
        connection.execute(
            "ALTER TABLE incoming_batches ADD COLUMN labeled_historical INTEGER NOT NULL DEFAULT 0"
        )
    incoming_ticket_columns = {row["name"] for row in connection.execute("PRAGMA table_info(incoming_tickets)")}
    if "low_information" not in incoming_ticket_columns:
        connection.execute(
            "ALTER TABLE incoming_tickets ADD COLUMN low_information INTEGER NOT NULL DEFAULT 0"
        )
    if "prediction_timestamp" not in incoming_ticket_columns:
        connection.execute("ALTER TABLE incoming_tickets ADD COLUMN prediction_timestamp TEXT")
    if "learning_priority" not in incoming_ticket_columns:
        connection.execute(
            "ALTER TABLE incoming_tickets ADD COLUMN learning_priority REAL NOT NULL DEFAULT 0"
        )
    if "needs_training_review" not in incoming_ticket_columns:
        connection.execute(
            "ALTER TABLE incoming_tickets ADD COLUMN needs_training_review INTEGER NOT NULL DEFAULT 0"
        )
    table = connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='tickets'").fetchone()
    if table:
        info = {row["name"]: dict(row) for row in connection.execute("PRAGMA table_info(tickets)")}
        if "user_name" not in info:
            connection.execute("ALTER TABLE tickets ADD COLUMN user_name TEXT")
            info = {row["name"]: dict(row) for row in connection.execute("PRAGMA table_info(tickets)")}
        needs_migration = info.get("request_id", {}).get("type", "").upper() != "TEXT" or "line1_participated" not in info
        if needs_migration:
            old_rows = connection.execute("SELECT * FROM tickets").fetchall()
            connection.execute("DROP INDEX IF EXISTS idx_tickets_filters")
            connection.execute("DROP INDEX IF EXISTS idx_tickets_description")
            connection.execute("ALTER TABLE tickets RENAME TO tickets_legacy")
            _create_tickets_table(connection)
            placeholders = ", ".join("?" for _ in _TICKET_COLUMNS)
            connection.executemany(
                f"INSERT INTO tickets ({', '.join(_TICKET_COLUMNS)}) VALUES ({placeholders})",
                [_legacy_values(row) for row in old_rows],
            )
            connection.execute("DROP TABLE tickets_legacy")
    else:
        _create_tickets_table(connection)
    connection.execute("INSERT OR REPLACE INTO metadata VALUES ('schema_version', ?)", (SCHEMA_VERSION,))
    connection.commit()
    connection.close()


def _record_values(
    row: dict[str, Any], headers: list[str], source_hash: str, updated_at: str, dataset_config: DatasetConfig
) -> tuple[Any, ...]:
    raw = {header: _json_value(row.get(header)) for header in headers}
    description = str(source_value(row, dataset_config, "description") or "")
    flag_values: list[int] = []
    for line in range(1, 5):
        reaction = parse_duration(source_value(row, dataset_config, f"line{line}_reaction", None)) or 0
        work = parse_duration(source_value(row, dataset_config, f"line{line}_work", None)) or 0
        flag_values.append(int(reaction > 0 or work > 0))
    flags = tuple(flag_values)
    overdue_value = source_value(row, dataset_config, "overdue", "")
    overdue = int(overdue_value is True or overdue_value == 1 or str(overdue_value).strip().lower() == "просрочен")
    return (
        _request_id(source_value(row, dataset_config, "request_id")),
        _json_value(source_value(row, dataset_config, "registration_date")),
        next((str(row.get(key) or "").strip() for key in ("Пользователь", "Заявитель", "ФИО пользователя", "user", "user_name") if str(row.get(key) or "").strip()), None),
        str(source_value(row, dataset_config, "service") or ""),
        str(source_value(row, dataset_config, "component") or "") or None,
        str(source_value(row, dataset_config, "category") or ""),
        str(source_value(row, dataset_config, "request_type") or ""),
        description,
        normalize_description(description),
        str(source_value(row, dataset_config, "criticality") or ""),
        str(source_value(row, dataset_config, "urgency") or ""),
        str(source_value(row, dataset_config, "priority") or ""),
        str(source_value(row, dataset_config, "service_class") or ""),
        str(source_value(row, dataset_config, "timezone") or ""),
        str(source_value(row, dataset_config, "status") or ""),
        overdue,
        str(source_value(row, dataset_config, "final_line") or ""),
        parse_duration(source_value(row, dataset_config, "actual_duration", None)),
        int(source_value(row, dataset_config, "clarifications_count", 0) or 0),
        parse_duration(source_value(row, dataset_config, "clarification_duration", None)),
        *flags,
        sum(flags),
        str(source_value(row, dataset_config, "result") or ""),
        json.dumps(raw, ensure_ascii=False, sort_keys=True),
        source_hash, updated_at,
    )


def _global_clarification_threshold(connection: sqlite3.Connection) -> int:
    values = sorted(row[0] for row in connection.execute("SELECT clarifications_count FROM tickets"))
    if not values:
        return 3
    return max(3, int(values[math.ceil(0.9 * (len(values) - 1))]))


def import_workbook(
    source: Path | None = None,
    database: Path = DATABASE_PATH,
    *,
    mode: Literal["snapshot", "delta", "upsert", "replace"] = "upsert",
    dataset_config: DatasetConfig | None = None,
) -> dict[str, Any]:
    if mode not in {"snapshot", "delta", "upsert", "replace"}:
        raise ValueError("Режим импорта должен быть snapshot, delta, upsert или replace")
    dataset_config = dataset_config or challenge_dataset_config()
    ensure_directories()
    source_path = find_dataset(source)
    headers, records = read_workbook(source_path, dataset_config)
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    ensure_database_schema(database)
    connection = _connect(database)
    if mode == "replace":
        connection.execute("DELETE FROM tickets")
    existing = {row["request_id"]: dict(row) for row in connection.execute(f"SELECT {', '.join(_TICKET_COLUMNS)} FROM tickets")}
    now = datetime.now(UTC).isoformat()
    inserted = updated = unchanged = 0
    compare_columns = [column for column in _TICKET_COLUMNS if column not in {"source_hash", "updated_at"}]
    for record in (_record_values(row, headers, source_hash, now, dataset_config) for row in records):
        item = dict(zip(_TICKET_COLUMNS, record, strict=True))
        previous = existing.get(item["request_id"])
        if previous is None:
            inserted += 1
        elif all(previous[column] == item[column] for column in compare_columns):
            unchanged += 1
            continue
        else:
            updated += 1
        update = ", ".join(f"{column}=excluded.{column}" for column in _TICKET_COLUMNS if column != "request_id")
        connection.execute(
            f"INSERT INTO tickets ({', '.join(_TICKET_COLUMNS)}) VALUES ({', '.join('?' for _ in _TICKET_COLUMNS)}) "
            f"ON CONFLICT(request_id) DO UPDATE SET {update}", record,
        )
    source_summary = inspect_workbook(source_path, dataset_config)
    database_count = connection.execute("SELECT COUNT(*) FROM tickets").fetchone()[0]
    threshold = _global_clarification_threshold(connection)
    result = {
        "status": "unchanged" if inserted == 0 and updated == 0 else "imported", "records": database_count,
        "inserted": inserted, "updated": updated, "unchanged": unchanged, "mode": mode, "database": str(database),
        "source_sha256": source_hash, "imported_at": now,
    }
    connection.commit()
    training_summary = training_corpus_summary(database, dataset_config)
    metadata = {
        "source_sha256": source_hash, "source_path": str(source_path),
        "dataset_sha256": training_summary["dataset_sha256"],
        "source_summary": json.dumps(source_summary, ensure_ascii=False),
        "dataset_summary": json.dumps(training_summary, ensure_ascii=False), "clarification_threshold": str(threshold),
        "last_import_summary": json.dumps(result, ensure_ascii=False), "imported_at": now,
    }
    connection.executemany("INSERT OR REPLACE INTO metadata VALUES (?, ?)", metadata.items())
    connection.execute(
        "INSERT INTO dataset_imports (source_hash, source_path, mode, imported_at, inserted, updated, unchanged, total) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        (source_hash, str(source_path), mode, now, inserted, updated, unchanged, database_count),
    )
    connection.commit()
    connection.close()
    return result | {
        "summary": source_summary | {"database_records": database_count},
        "training_summary": training_summary,
        "clarification_threshold": threshold,
    }


def load_rows(database: Path = DATABASE_PATH) -> list[dict[str, Any]]:
    ensure_database_schema(database)
    connection = _connect(database)
    rows = [dict(row) for row in connection.execute("SELECT * FROM tickets ORDER BY request_id")]
    connection.close()
    return rows


def metadata_value(key: str, database: Path = DATABASE_PATH) -> Any:
    ensure_database_schema(database)
    connection = _connect(database)
    row = connection.execute("SELECT value FROM metadata WHERE key=?", (key,)).fetchone()
    connection.close()
    if row is None:
        return None
    try:
        return json.loads(row[0])
    except json.JSONDecodeError:
        return row[0]
