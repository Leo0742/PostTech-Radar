from __future__ import annotations

import csv
import sqlite3
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from statistics import median
from typing import Any

from openpyxl import load_workbook

from app.core.config import DATABASE_PATH
from app.services.data_service import ensure_database_schema

_REQUIRED_EVENT_COLUMNS = ("request_id", "status", "entered_at")


def _event_rows(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
    elif path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() for value in next(values)]
        rows = [dict(zip(headers, row, strict=True)) for row in values if any(value is not None for value in row)]
        workbook.close()
    else:
        raise ValueError("История статусов должна быть CSV или XLSX")
    missing = [column for column in _REQUIRED_EVENT_COLUMNS if column not in (rows[0] if rows else {})]
    if missing:
        raise ValueError(f"В истории статусов отсутствуют колонки: {', '.join(missing)}")
    return rows


def _iso(value: Any, *, optional: bool = False) -> str | None:
    if value is None or str(value).strip() == "":
        if optional:
            return None
        raise ValueError("entered_at не может быть пустым")
    if isinstance(value, datetime):
        return value.isoformat()
    return datetime.fromisoformat(str(value).strip()).isoformat()


def import_status_history(path: Path, database: Path = DATABASE_PATH) -> dict[str, Any]:
    ensure_database_schema(database)
    rows = _event_rows(path)
    connection = sqlite3.connect(database)
    inserted = 0
    for row in rows:
        cursor = connection.execute(
            "INSERT OR IGNORE INTO status_events (request_id, status, entered_at, left_at, source) VALUES (?, ?, ?, ?, ?)",
            (str(row["request_id"]).strip(), str(row["status"]).strip(), _iso(row["entered_at"]),
             _iso(row.get("left_at"), optional=True), str(row.get("source") or path.name)),
        )
        inserted += cursor.rowcount
    connection.commit()
    connection.close()
    return {"inserted": inserted, "total": len(rows), "source": str(path)}


def _status_history(connection: sqlite3.Connection) -> dict[str, Any]:
    rows = connection.execute(
        "SELECT request_id, status, entered_at, left_at, source FROM status_events ORDER BY request_id, entered_at"
    ).fetchall()
    if not rows:
        return {
            "available": False,
            "message": "В предоставленной выгрузке отсутствует журнал изменения статусов. Истинный граф переходов между статусами невозможно восстановить без событийной истории.",
            "required_columns": [*_REQUIRED_EVENT_COLUMNS, "left_at", "source"], "transitions": [], "dwell": [],
        }
    per_request: dict[str, list[sqlite3.Row]] = defaultdict(list)
    for row in rows:
        per_request[row["request_id"]].append(row)
    transitions: Counter[tuple[str, str]] = Counter()
    dwell: dict[str, list[float]] = defaultdict(list)
    for events in per_request.values():
        for index, event in enumerate(events):
            if index + 1 < len(events):
                transitions[(event["status"], events[index + 1]["status"])] += 1
            if event["left_at"]:
                seconds = (datetime.fromisoformat(event["left_at"]) - datetime.fromisoformat(event["entered_at"])).total_seconds()
                if seconds >= 0:
                    dwell[event["status"]].append(seconds)
    total = sum(transitions.values())
    return {
        "available": True, "message": "Граф построен только по импортированному журналу событий.",
        "required_columns": [*_REQUIRED_EVENT_COLUMNS, "left_at", "source"], "event_count": len(rows),
        "request_count": len(per_request),
        "transitions": [{"from": source, "to": target, "count": count, "share": round(count / total, 4) if total else 0}
                        for (source, target), count in transitions.most_common()],
        "dwell": [{"status": status, "median_seconds": int(median(values)), "sample_n": len(values)}
                  for status, values in sorted(dwell.items())],
    }


def _participation(connection: sqlite3.Connection) -> dict[str, Any]:
    rows = connection.execute(
        "SELECT line1_participated, line2_participated, line3_participated, line4_participated, "
        "participants_count, final_line, overdue, actual_duration_seconds, category FROM tickets"
    ).fetchall()
    combinations: dict[str, list[sqlite3.Row]] = defaultdict(list)
    resolver_edges: Counter[tuple[str, str]] = Counter()
    co_participation: Counter[tuple[str, str]] = Counter()
    for row in rows:
        lines = [str(index) for index in range(1, 5) if row[f"line{index}_participated"]]
        combinations["+".join(lines) if lines else "Нет данных"].append(row)
        resolver = str(row["final_line"]).replace("(", "").replace(" линия)", "")
        for line in lines:
            resolver_edges[(line, resolver)] += 1
        for left_index, left in enumerate(lines):
            for right in lines[left_index + 1:]:
                co_participation[(left, right)] += 1
    combo_rows = []
    for combination, group in combinations.items():
        durations = [row["actual_duration_seconds"] for row in group if row["actual_duration_seconds"] is not None]
        overdue = sum(row["overdue"] for row in group)
        combo_rows.append({"combination": combination, "count": len(group),
                           "overdue_rate": round(overdue / len(group), 4),
                           "median_duration_seconds": int(median(durations)) if durations else 0})
    single_durations = sorted(
        int(row["actual_duration_seconds"])
        for row in rows
        if row["participants_count"] <= 1 and row["actual_duration_seconds"] is not None
    )
    multi_durations = sorted(
        int(row["actual_duration_seconds"])
        for row in rows
        if row["participants_count"] > 1 and row["actual_duration_seconds"] is not None
    )

    def percentile(values: list[int], share: float) -> int:
        if not values:
            return 0
        index = min(len(values) - 1, max(0, round((len(values) - 1) * share)))
        return values[index]

    return {
        "edge_semantics": "участие → финальная линия; не хронологический переход",
        "combinations": sorted(combo_rows, key=lambda item: item["count"], reverse=True),
        "resolver_edges": [{"participant": participant, "resolver": resolver, "count": count}
                           for (participant, resolver), count in resolver_edges.most_common()],
        "co_participation": [{"line_a": left, "line_b": right, "count": count}
                             for (left, right), count in co_participation.most_common()],
        "duration_groups": {
            "single_line": {
                "sample_n": len(single_durations),
                "median_seconds": int(median(single_durations)) if single_durations else 0,
                "p90_seconds": percentile(single_durations, 0.9),
            },
            "multi_line": {
                "sample_n": len(multi_durations),
                "median_seconds": int(median(multi_durations)) if multi_durations else 0,
                "p90_seconds": percentile(multi_durations, 0.9),
            },
        },
        "multi_line_count": sum(row["participants_count"] > 1 for row in rows), "ticket_count": len(rows),
    }


def process_summary(*, database: Path = DATABASE_PATH) -> dict[str, Any]:
    ensure_database_schema(database)
    connection = sqlite3.connect(database)
    connection.row_factory = sqlite3.Row
    result = {"status_history": _status_history(connection), "participation": _participation(connection)}
    connection.close()
    return result
